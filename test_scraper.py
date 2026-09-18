"""Offline behavioral tests; no network, Chrome or credentials required."""

from contextlib import redirect_stderr, redirect_stdout
from dataclasses import asdict, replace
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock, patch

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from webscraper.cli import main
from webscraper.config import Config, normalize_url, same_origin
from webscraper.engine import Crawler
from webscraper.exporters import export_all, render_report
from webscraper.extraction import parse_detail, parse_listing
from webscraper.sources import FetchError, FixtureSource, Page, PageTimeout, SeleniumSource
from webscraper.storage import output_lock, read_checkpoint, save_checkpoint

ROOT = Path(__file__).resolve().parent.parent
NOW = "2026-09-18T10:00:00+00:00"


class Base(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.output = Path(self.temp.name)
        self.config = Config.load(ROOT / "examples/demo.json")
        self.source = FixtureSource(ROOT / "examples/site")

    def crawler(self, source=None, **kwargs):
        return Crawler(self.config, source or self.source, limit=kwargs.pop("limit", 15),
                       checkpoint=self.output / "checkpoint.json", fingerprint=self.config.fingerprint("demo"),
                       sleep=kwargs.pop("sleep", lambda _: None), clock=lambda: NOW, mode="demo", **kwargs)


class ConfigurationTests(Base):
    def test_invalid_settings_fail_before_navigation(self):
        changes = [
            {"listing_url": "javascript:alert(1)"}, {"listing_url": "https://user:pass@test.test/"},
            {"link_selector": "["}, {"request_delay": 0}, {"timeout": float("nan")},
            {"retries": True}, {"retries": 1.5}, {"max_listing_pages": 0},
            {"required_fields": ["not_configured"]}, {"required_fields": []},
            {"fields": {"status": "h1"}}, {"next_selector": ""}, {"unknown": 12},
        ]
        for change in changes:
            with self.subTest(change=change), self.assertRaises(ValueError):
                Config.from_dict({**asdict(self.config), **change})

    def test_url_resolution_keeps_queries_and_strips_fragments(self):
        self.assertEqual(normalize_url("../record?id=2#part", "https://DEMO.test:443/list/page"),
                         "https://demo.test/record?id=2")
        self.assertTrue(same_origin("https://demo.test:443/x", "https://demo.test/y"))
        self.assertFalse(same_origin("http://demo.test/x", "https://demo.test/"))
        self.assertFalse(same_origin("https://demo.test.evil.test/", "https://demo.test/"))

    def test_disallowed_links_are_rejected(self):
        for link in ["#top", "", "mailto:a@test.test", "javascript:alert(1)", "file:///tmp/x"]:
            with self.subTest(link=link):
                self.assertIsNone(normalize_url(link, self.config.listing_url))

    def test_invalid_limit(self):
        for limit in [0, True, 5001]:
            with self.subTest(limit=limit), self.assertRaises(ValueError):
                self.crawler(limit=limit)


class ExtractionTests(Base):
    def test_duplicate_labels_and_all_blocks_survive(self):
        url = "https://demo.test/records/atlas.html"
        page = self.source.fetch(url, "body")
        record = parse_detail(page.html, url, url, self.config, NOW)
        self.assertEqual(record.status, "ok")
        self.assertEqual([b["label"] for b in record.blocks], ["Service", "Service", "Service_2", "Hours"])
        self.assertEqual(record.blocks[1]["value"], "Editorial design")

    def test_missing_required_field_is_partial_not_success(self):
        record = parse_detail('<main class="main-content"><h1></h1></main>', "u", "u", self.config, NOW)
        self.assertEqual(record.status, "partial")
        self.assertIn("Required field is empty: name", record.issues)

    def test_missing_optional_field_is_empty_not_invented(self):
        url = "https://demo.test/records/northstar.html"
        record = parse_detail(self.source.fetch(url, "body").html, url, url, self.config, NOW)
        self.assertEqual(record.status, "ok")
        self.assertEqual(record.fields["email"], "")

    def test_incomplete_block_is_retained_and_flagged(self):
        html = '<h1 class="record-name">Name</h1><b class="category-tag">Test</b><div class="data-block"><b class="block-label">Service</b></div>'
        record = parse_detail(html, "u", "u", self.config, NOW)
        self.assertEqual(record.blocks, [{"label": "Service", "value": ""}])
        self.assertEqual(record.status, "partial")

    def test_listing_deduplicates_and_excludes_external_links(self):
        url = self.config.listing_url
        links, next_url = parse_listing(self.source.fetch(url, "body").html, url, self.config)
        self.assertEqual(links, ["https://demo.test/records/atlas.html", "https://demo.test/records/cedar.html"])
        self.assertEqual(next_url, "https://demo.test/directory/page-2.html")

    def test_empty_listing_is_an_error(self):
        with self.assertRaisesRegex(ValueError, "No detail links"):
            parse_listing("<h1>No matches</h1>", self.config.listing_url, self.config)

    def test_external_next_link_is_an_error(self):
        with self.assertRaisesRegex(ValueError, "next-page"):
            parse_listing('<a class="record-link" href="/x">X</a><a class="next-page" href="https://outside.test/">Next</a>', self.config.listing_url, self.config)


class PipelineTests(Base):
    def test_end_to_end_demo_visits_two_pages_once_per_record(self):
        records, summary = self.crawler().run()
        self.assertEqual(len(records), 3)
        self.assertEqual(summary["counts"], {"ok": 3, "partial": 0, "error": 0})
        self.assertEqual(summary["listing_pages"], 2)
        self.assertEqual(summary["page_attempts"], 5)
        self.assertEqual(sum(len(r.blocks) for r in records), 10)
        self.assertEqual(summary["stop_reason"], "end_of_listing")

    def test_sample_is_bounded(self):
        records, summary = self.crawler(limit=1).run()
        self.assertEqual(len(records), 1)
        self.assertEqual(summary["page_attempts"], 2)
        self.assertEqual(summary["stop_reason"], "record_limit")
        self.assertTrue(summary["warnings"])

    def test_listing_page_limit_is_visible(self):
        self.config = replace(self.config, max_listing_pages=1)
        records, summary = self.crawler().run()
        self.assertEqual(len(records), 2)
        self.assertEqual(summary["stop_reason"], "listing_page_limit")

    def test_pagination_cycle_fails_instead_of_looping(self):
        html = '<a class="record-link" href="/record">R</a><a class="next-page" href="index.html">Next</a>'
        source = Mock()
        source.fetch.return_value = Page(self.config.listing_url, html)
        with self.assertRaisesRegex(FetchError, "loop"):
            self.crawler(source).run()
        self.assertEqual(source.fetch.call_count, 1)

    def test_timeout_retries_are_bounded_and_paced(self):
        source, sleep = Mock(), Mock()
        source.fetch.side_effect = [PageTimeout("wait"), Page(self.config.listing_url, "<p>ok</p>")]
        crawler = self.crawler(source, sleep=sleep)
        _, attempts = crawler.fetch(self.config.listing_url, "body")
        self.assertEqual(attempts, 2)
        sleep.assert_called_once_with(3.0)

    def test_non_timeout_error_is_not_retried(self):
        source = Mock()
        source.fetch.side_effect = FetchError("Denied")
        with self.assertRaises(FetchError):
            self.crawler(source).fetch(self.config.listing_url, "body")
        self.assertEqual(source.fetch.call_count, 1)

    def test_detail_failure_is_saved_and_other_records_continue(self):
        fixture = self.source
        source = Mock()
        def fetch(url, selector):
            if url.endswith("atlas.html"):
                raise PageTimeout("wait")
            return fixture.fetch(url, selector)
        source.fetch.side_effect = fetch
        records, summary = self.crawler(source).run()
        self.assertEqual(summary["counts"], {"ok": 2, "partial": 0, "error": 1})
        self.assertEqual(records[0].attempts, 2)
        self.assertEqual(len(read_checkpoint(self.output / "checkpoint.json", self.config.fingerprint("demo"))), 3)

    def test_interruption_preserves_completed_records(self):
        fixture = self.source
        source = Mock()
        def fetch(url, selector):
            if url.endswith("cedar.html"):
                raise KeyboardInterrupt()
            return fixture.fetch(url, selector)
        source.fetch.side_effect = fetch
        with self.assertRaises(KeyboardInterrupt):
            self.crawler(source).run()
        previous = read_checkpoint(self.output / "checkpoint.json", self.config.fingerprint("demo"))
        self.assertEqual(len(previous), 1)
        records, summary = self.crawler(previous=previous).run()
        self.assertEqual(summary["reused_records"], 1)
        self.assertEqual(summary["page_attempts"], 4)
        self.assertEqual(len(records), 3)

    def test_resume_reuses_only_successful_records(self):
        records, _ = self.crawler().run()
        records[0].status = "partial"
        records[1].status = "error"
        previous = {r.requested_url: r for r in records}
        result, summary = self.crawler(previous=previous).run()
        self.assertEqual(summary["reused_records"], 1)
        self.assertEqual(summary["attempted_records"], 2)
        self.assertTrue(all(r.status == "ok" for r in result))

    def test_smaller_resume_keeps_durable_records_outside_sample(self):
        records, _ = self.crawler().run()
        self.crawler(limit=1, previous={r.requested_url: r for r in records}).run()
        durable = read_checkpoint(self.output / "checkpoint.json", self.config.fingerprint("demo"))
        self.assertEqual(len(durable), 3)

    def test_off_origin_redirect_fails(self):
        source = Mock()
        source.fetch.return_value = Page("https://outside.test/", "<h1>Wrong origin</h1>")
        with self.assertRaisesRegex(FetchError, "origin"):
            self.crawler(source).run()


class PersistenceAndExportTests(Base):
    def test_changed_config_rejects_checkpoint(self):
        self.crawler().run()
        with self.assertRaisesRegex(ValueError, "differs"):
            read_checkpoint(self.output / "checkpoint.json", "a-different-fingerprint")

    def test_truncated_checkpoint_is_not_silently_overwritten(self):
        path = self.output / "checkpoint.json"
        path.write_text('{"records":', encoding="utf-8")
        with self.assertRaises(ValueError):
            read_checkpoint(path, "x")
        self.assertEqual(path.read_text(), '{"records":')

    def test_malformed_checkpoint_record_is_rejected(self):
        records, _ = self.crawler().run()
        records[0].blocks = ["not a data block"]
        save_checkpoint(self.output / "checkpoint.json", self.config.fingerprint("demo"), records)
        with self.assertRaisesRegex(ValueError, "invalid record"):
            read_checkpoint(self.output / "checkpoint.json", self.config.fingerprint("demo"))

    def test_same_output_directory_cannot_run_concurrently(self):
        with output_lock(self.output):
            with self.assertRaisesRegex(ValueError, "locked"):
                with output_lock(self.output):
                    self.fail("Lock allowed a second writer")
        self.assertFalse((self.output / ".run.lock").exists())

    def test_json_preserves_full_data_and_excel_strings_are_literal(self):
        records, summary = self.crawler().run()
        records[0].fields["name"] = '=HYPERLINK("https://example.test", "test")'
        records[0].fields["phone"] = "001234"
        records[0].blocks.append({"label": "Formula-like", "value": "=1+1"})
        export_all(self.output, records, list(self.config.fields), summary)
        data = json.loads((self.output / "records.json").read_text())
        self.assertEqual(data[0]["fields"]["phone"], "001234")
        book = load_workbook(self.output / "records.xlsx")
        try:
            self.assertEqual(book["Records"]["A2"].data_type, "s")
            self.assertEqual(book["Records"]["B2"].value, "001234")
            self.assertEqual(book["Data blocks"].max_row, 12)
            self.assertFalse(any(c.data_type == "f" for sheet in book for row in sheet for c in row))
        finally:
            book.close()

    def test_excel_limit_is_reported_json_remains_complete(self):
        records, summary = self.crawler().run()
        records[0].fields["notes"] = "a" * 33000 + "\x01"
        export_all(self.output, records, list(self.config.fields), summary)
        data = json.loads((self.output / "records.json").read_text())
        self.assertEqual(len(data[0]["fields"]["notes"]), 33001)
        self.assertIn("JSON retains full values", summary["warnings"][-1])

    def test_report_escapes_untrusted_html(self):
        records, summary = self.crawler().run()
        records[0].fields["name"] = '<img src=x onerror="alert(1)">'
        records[0].blocks[0]["value"] = "</script><script>alert(2)</script>"
        html = render_report(records, summary)
        soup = BeautifulSoup(html, "html.parser")
        self.assertEqual(len(soup.find_all("script")), 1)
        self.assertIsNone(soup.find("img"))
        self.assertIn('<img src=x onerror="alert(1)">', soup.get_text())


class CliAndBrowserAdapterTests(Base):
    def call_cli(self, *args):
        with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            return main(list(args))

    def test_demo_cli_and_resume_never_launch_chrome(self):
        with patch("webscraper.cli.SeleniumSource") as chrome:
            self.assertEqual(self.call_cli("--demo", "--output", str(self.output)), 0)
            self.assertEqual(self.call_cli("--demo", "--output", str(self.output), "--resume"), 0)
            chrome.assert_not_called()
        summary = json.loads((self.output / "run_report.json").read_text())
        self.assertEqual(summary["reused_records"], 3)
        self.assertEqual(summary["page_attempts"], 2)

    def test_existing_output_requires_explicit_action(self):
        (self.output / "records.json").write_text("keep me")
        self.assertEqual(self.call_cli("--demo", "--output", str(self.output)), 2)
        self.assertEqual((self.output / "records.json").read_text(), "keep me")

    def test_overwrite_keeps_unrelated_files(self):
        (self.output / "unrelated.txt").write_text("keep")
        (self.output / "records.json").write_text("old")
        self.assertEqual(self.call_cli("--demo", "--output", str(self.output), "--overwrite"), 0)
        self.assertEqual((self.output / "unrelated.txt").read_text(), "keep")

    def test_missing_resume_checkpoint_is_an_error(self):
        self.assertEqual(self.call_cli("--demo", "--output", str(self.output), "--resume"), 2)

    def test_partial_records_have_nonzero_exit_code(self):
        page = Page("https://demo.test/records/missing.html", '<main class="main-content"></main>')
        source = Mock()
        source.fetch.side_effect = [Page(self.config.listing_url, '<a class="record-link" href="/records/missing.html">R</a>'), page]
        with patch("webscraper.cli.FixtureSource", return_value=source):
            self.assertEqual(self.call_cli("--demo", "--output", str(self.output)), 1)
        source.close.assert_called_once()

    def test_browser_adapter_waits_and_always_closes(self):
        with patch("selenium.webdriver.Chrome") as chrome, patch("selenium.webdriver.support.ui.WebDriverWait") as wait:
            driver = chrome.return_value
            driver.current_url = self.config.listing_url
            driver.page_source = "<main>Ready</main>"
            source = SeleniumSource(self.config)
            page = source.fetch(self.config.listing_url, ".directory")
            self.assertEqual(page.html, "<main>Ready</main>")
            wait.return_value.until.assert_called_once()
            source.close()
            driver.quit.assert_called_once()

    def test_browser_timeout_becomes_retryable_error(self):
        from selenium.common.exceptions import TimeoutException
        with patch("selenium.webdriver.Chrome") as chrome:
            chrome.return_value.get.side_effect = TimeoutException("wait")
            source = SeleniumSource(self.config)
            with self.assertRaises(PageTimeout):
                source.fetch(self.config.listing_url, ".directory")
            source.close()


if __name__ == "__main__":
    unittest.main()

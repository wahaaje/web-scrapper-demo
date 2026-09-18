"""JSON, reviewable HTML and Excel exports with literal spreadsheet strings."""

from dataclasses import asdict
from html import escape
import json
import os
from pathlib import Path
import re
from string import Template
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .extraction import Record
from .storage import atomic_text

INVALID_XML = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F\uFFFE\uFFFF]")


def save_excel(path: Path, records: list[Record], fields: list[str], summary: dict) -> list[str]:
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Run summary"
    overview.append(["Metric", "Value"])
    for key in ("mode", "started_at", "finished_at", "records", "listing_pages",
                "page_attempts", "reused_records", "stop_reason"):
        overview.append([key, summary[key]])
    for key, value in summary["counts"].items():
        overview.append([f"status_{key}", value])
    for warning in summary["warnings"]:
        overview.append(["warning", warning])
    sheet = workbook.create_sheet("Records")
    headers = [*fields, "status", "issues", "attempts", "source_url", "requested_url", "scraped_at"]
    sheet.append(headers)
    blocks = workbook.create_sheet("Data blocks")
    blocks.append(["requested_url", "block_index", "label", "value"])
    warnings: list[str] = []
    changed = 0

    def literal_row(target, values):
        nonlocal changed
        row = target.max_row + 1
        for column, value in enumerate(values, start=1):
            cell = target.cell(row, column)
            if isinstance(value, str):
                cleaned = INVALID_XML.sub("\uFFFD", value)
                if cleaned != value or len(cleaned) > 32767:
                    changed += 1
                cell.value = cleaned[:32767]
                # Treat every extracted string literally, including leading '='.
                cell.data_type = "s"
            else:
                cell.value = value

    for record in records:
        literal_row(sheet, [*(record.fields.get(f, "") for f in fields), record.status,
                            " | ".join(record.issues), record.attempts, record.source_url,
                            record.requested_url, record.scraped_at])
        for number, block in enumerate(record.blocks, start=1):
            literal_row(blocks, [record.requested_url, number, block["label"], block["value"]])
    if changed:
        warnings.append(f"Excel adjusted {changed} cell(s) for XML/length limits; JSON retains full values.")
        literal_row(overview, ["export_warning", warnings[-1]])
    palette = {"ok": "DCF5E9", "partial": "FFF1CE", "error": "FFE1DF"}
    status_column = headers.index("status") + 1
    for row in range(2, sheet.max_row + 1):
        cell = sheet.cell(row, status_column)
        cell.fill = PatternFill("solid", fgColor=palette[cell.value])
    for target in workbook:
        target.freeze_panes = "A2"
        target.auto_filter.ref = target.dimensions
        target.sheet_view.showGridLines = False
        target.row_dimensions[1].height = 30
        for cell in target[1]:
            cell.fill = PatternFill("solid", fgColor="102D38")
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
            cell.alignment = Alignment(vertical="center")
        for row in target.iter_rows(min_row=2):
            target.row_dimensions[row[0].row].height = 46
            for cell in row:
                cell.font = Font(name="Calibri", size=11, color="102D38")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in range(1, target.max_column + 1):
            header = str(target.cell(1, column).value)
            width = 48 if header in {"value", "Value", "issues", "source_url", "requested_url"} else 26
            if header in {"status", "attempts", "block_index"}:
                width = 14
            target.column_dimensions[get_column_letter(column)].width = width
    fd, temp = tempfile.mkstemp(suffix=".xlsx", dir=path.parent)
    os.close(fd)
    try:
        workbook.save(temp)
        os.replace(temp, path)
    finally:
        workbook.close()
        if os.path.exists(temp):
            os.unlink(temp)
    return warnings


def render_report(records: list[Record], summary: dict) -> str:
    cards = []
    for record in records:
        field_rows = "".join(
            f"<div><dt>{escape(key.replace('_', ' ').title())}</dt><dd>{escape(value) or '—'}</dd></div>"
            for key, value in record.fields.items()
        )
        block_rows = "".join(
            f"<tr><td>{i}</td><td>{escape(b['label'])}</td><td>{escape(b['value'])}</td></tr>"
            for i, b in enumerate(record.blocks, start=1)
        )
        issues = "".join(f"<li>{escape(issue)}</li>" for issue in record.issues)
        source = escape(record.source_url or record.requested_url)
        name = escape(record.fields.get("name") or "Unnamed record")
        status = escape(record.status)
        cards.append(f'''<article class="record" data-status="{status}">
          <div class="record-head"><h2>{name}</h2><span class="pill {status}">{status}</span></div>
          <dl>{field_rows}</dl>
          <details><summary>{len(record.blocks)} data blocks · inspect extracted sections</summary>
            <div class="table-scroll"><table><thead><tr><th>#</th><th>Label</th><th>Value</th></tr></thead>
            <tbody>{block_rows}</tbody></table></div></details>
          {('<ul class="issues">' + issues + '</ul>') if issues else ''}
          <footer><span>Source · {source}</span><span>{record.attempts} attempt(s) · {escape(record.scraped_at)}</span></footer>
        </article>''')
    template = Template(Path(__file__).with_name("report.html").read_text(encoding="utf-8"))
    counts = summary["counts"]
    note = "Local fixture demo · all organizations and contact details are fictional." if summary["mode"] == "demo" else "Browser extraction · review results before using them downstream."
    warnings = "".join(f"<li>{escape(w)}</li>" for w in summary["warnings"])
    return template.substitute(
        records="".join(cards), count=summary["records"], ok=counts["ok"],
        partial=counts["partial"], errors=counts["error"], pages=summary["listing_pages"],
        date=escape(summary["finished_at"]), mode=escape(summary["mode"].upper()),
        note=note, warnings=f'<ul class="notice">{warnings}</ul>' if warnings else "",
        reused=summary["reused_records"], stop=escape(summary["stop_reason"].replace("_", " ")),
    )


def export_all(output: Path, records: list[Record], fields: list[str], summary: dict) -> None:
    # JSON is the canonical complete dataset, including repeated and long data blocks.
    atomic_text(output / "records.json", json.dumps([asdict(r) for r in records], indent=2, ensure_ascii=False))
    summary["warnings"].extend(save_excel(output / "records.xlsx", records, fields, summary))
    atomic_text(output / "report.html", render_report(records, summary))
    atomic_text(output / "run_report.json", json.dumps(summary, indent=2, ensure_ascii=False))
